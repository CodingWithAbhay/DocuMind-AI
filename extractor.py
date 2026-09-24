"""Text extraction for document types (PDF, Word, PowerPoint, Excel, Plain Text, Markdown)."""

from __future__ import annotations

import io


class UnsupportedDocument(Exception):
    pass


SUPPORTED_EXTENSIONS = ["pdf", "docx", "pptx", "xlsx", "xls", "txt", "md"]


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def _docx(data: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    parts = []
    for i, slide in enumerate(prs.slides, start=1):
        parts.append(f"--- Slide {i} ---")
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
    return "\n".join(parts)


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"--- Sheet: {ws.title} ---")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _xls(data: bytes) -> str:
    import xlrd

    workbook = xlrd.open_workbook(file_contents=data)
    parts = []
    for sheet in workbook.sheets():
        parts.append(f"--- Sheet: {sheet.name} ---")
        for row_index in range(sheet.nrows):
            cells = [str(value) for value in sheet.row_values(row_index) if value != ""]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_text(filename: str, data: bytes) -> str:
    """Return plain text for a supported document, else raise UnsupportedDocument."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    handlers = {
        "pdf": _pdf,
        "docx": _docx,
        "pptx": _pptx,
        "xlsx": _xlsx,
        "xls": _xls,
    }

    if ext in handlers:
        text = handlers[ext](data)
    elif ext in ("txt", "md"):
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError:
            raise UnsupportedDocument(
                f"'{filename}' is not a readable text document."
            ) from None
    else:
        raise UnsupportedDocument(f"'{filename}' has an unsupported file format.")

    text = text.replace("\r", "").strip()
    if not text:
        raise UnsupportedDocument(f"No readable text was found in '{filename}'.")
    return text